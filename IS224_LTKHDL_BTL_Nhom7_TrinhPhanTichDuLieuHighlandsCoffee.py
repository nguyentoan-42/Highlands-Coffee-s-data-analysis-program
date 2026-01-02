import pandas as pd
import numpy as np
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
import openpyxl
import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg



# HÀM LÀM SẠCH DỮ LIỆU
def clean_highlands_data(df):
    # Xử lí dòng trống và dòng trùng lặp
    df = df.dropna(how='all')
    df = df.drop_duplicates(keep='first')

    df = df.reset_index(drop=True)


    # Store_name
    df["store_name"] = df["store_name"].str.title().str.strip()


    # City
    def normalize_city(city):
        if pd.isna(city):
            return "Unknown"
        city = city.lower().strip()
        if city in ["hcm", "ho chi minh", "hồ chí minh", "hcmc", "hochiminh"]:
            return "HCM"
        if city in ["ha noi", "hanoi", "hà nội"]:
            return "Hà Nội"
        if city in ["da nang", "danang", "đà nẵng"]:
            return "Đà Nẵng"
        return city.title()
    
    df["city"] = df["city"].apply(normalize_city)


    # Revenue
        # R - Clean
    def clean_revenue(value):
        try:
            value = str(value).replace(",", "").split(".")[0]

            if int(value) < 50000:
                return 0

            return int(value)
        except:
            return np.nan
        
    df["revenue"] = df["revenue"].apply(clean_revenue)

        # R - Missing value
    group_median_revenue = df.groupby("city")["revenue"].transform("median")
    global_median_revenue = df["revenue"].median()
    df["revenue"] = df["revenue"].fillna(group_median_revenue).fillna(global_median_revenue).astype(int)


    # Customers
    df["customers"] = pd.to_numeric(df["customers"], errors="coerce")

    group_median_customers = df.groupby("city")["customers"].transform("median")
    global_median_customers = df["customers"].median()
    df["customers"] = df["customers"].fillna(group_median_customers).fillna(global_median_customers).astype(int)


    # Month
    df["month"] = pd.to_datetime(df["month"], format="%m-%Y", errors="coerce")

    df = df.sort_values("month", ascending=False).reset_index(drop=True)
    df["month"] = df["month"].dt.strftime("%m-%Y")


    # Profit
    df["profit"] = pd.to_numeric(df["profit"], errors="coerce")
    df["profit"] = df["profit"].fillna(0).astype(int)


    return df



# DATAFRAME TOÀN CỤC
df_raw = None
df_clean = None



# CHỌN FILE
def select_file():
    global df_raw, df_clean

    file_path = filedialog.askopenfilename(
        title="Chọn file CSV hoặc Excel",
        filetypes=[
            ("All supported", "*.csv;*.xlsx;*.xls"),
            ("CSV files", "*.csv"),
            ("Excel files", "*.xlsx;*.xls"),
            ("All files", "*.*")
        ]
    )

    if not file_path:
        return
    
    try:
        if file_path.lower().endswith(".csv"):
            df_raw = pd.read_csv(file_path, encoding="utf-8-sig")
        
        elif file_path.lower().endswith((".xlsx", ".xls")):
            df_raw = pd.read_excel(file_path)

        else:
            messagebox.showerror("Lỗi", "Định dạng file không được hỗ trợ!\nChỉ hỗ trợ: .csv, .xlsx, .xls")
            return

        
        if df_raw.empty:
            messagebox.showerror("Lỗi", "File rỗng!")
            df_raw = None
            return


        required_cols = ["store_id", "store_name", "city", "revenue", "customers", "month", "profit"]
        missing_cols = set(required_cols) - set(df_raw.columns)

        if missing_cols:
            messagebox.showerror("Lỗi", f"File thiếu cột:\n{", ".join(missing_cols)}")
            df_raw = None
            return
        

        df_clean = None


        show_table(df_raw)
        filename = file_path.split("/")[-1]
        input_label.config(text=f"{filename} ({len(df_raw)} dòng)")
        status_label.config(text="Đã nhận dữ liệu gốc")


        clean_btn.config(state="normal")
        save_btn.config(state="disabled")
        id_sort_btn.config(state="normal")
        city_sort_btn.config(state="normal")
        search_btn.config(state="normal")
        delete_btn.config(state="normal")


    except Exception as e:
        messagebox.showerror("Lỗi", f"Không thể đọc được file:\n{str(e)}")
        df_raw = None



# LÀM SẠCH DỮ LIỆU
def clean_data():
    global df_clean


    if df_raw is None:
        messagebox.showwarning("Chưa có dữ liệu", "Vui lòng chọn file trước")
        return
    

    try:
        df_clean = clean_highlands_data(df_raw.copy())
        show_table(df_clean)
        messagebox.showinfo("Làm sạch dữ liệu", "Làm sạch dữ liệu hoàn tất!")


        status_label.config(text="Dữ liệu đã được làm sạch")
        save_btn.config(state="normal")
        analyze_btn.config(state="normal")
        chart_btn.config(state="normal")

    except Exception as e:
        messagebox.showerror("Lỗi", f"Lỗi khi làm sạch dữ liệu:\n{str(e)}")



# LƯU FILE
def save_file():
    if df_clean is None:
        messagebox.showwarning("Chưa có dữ liệu", "Vui lòng làm sạch dữ liệu trước")
        return

    
    save_path = filedialog.asksaveasfilename(
        title="Chọn nơi lưu file sau khi đã làm sạch",
        defaultextension=".csv",
        filetypes=[
            ("CSV files", "*.csv"),
            ("Excel files", "*.xlsx;*.xls"),
            ("All files", "*.*")
        ]
    )


    if save_path:
        try:
            df_clean.to_csv(save_path, index=False, encoding="utf-8-sig")
            messagebox.showinfo("Hoàn tất", f"Đã lưu file:\n{save_path}")
            status_label.config(text=f"Đã lưu: {save_path.split("/")[-1]}")
        
        except Exception as e:
            messagebox.showerror("Lỗi", f"Không thể lưu file:\n{str(e)}")



# HIỂN THỊ DỮ LIỆU TRONG BẢNG
def show_table(df):
    table.delete(*table.get_children())


    table["columns"] = list(df.columns)
    table["show"] = "headings"


    for col in df.columns:
        table.heading(col, text=col)
        table.column(col, anchor="center")


    for idx, row in df.iterrows():
        if idx >= 1000:
            break
        table.insert("", "end", values=list(row))


    if len(df) > 1000:
        table_info.config(text=f"Hiển thị 1000/{len(df)} dòng đầu tiên")

    else:
        table_info.config(text=f"Hiển thị {len(df)} dòng")



# SẮP XẾP THEO ID CỬA HÀNG
def sort_by_storeid():
    global df_raw, df_clean

    target_df = df_clean if df_clean is not None else df_raw

    if target_df is not None:
        target_df = target_df.sort_values(by="store_id").reset_index(drop=True)

        if df_clean is not None:
            df_clean = target_df

        else:
            df_raw = target_df

        show_table(target_df)
        status_label.config(text="Đã sắp xếp theo ID cửa hàng")

    else:
        messagebox.showwarning("Thông báo", "Vui lòng tải dữ liệu lên trước khi sắp xếp")



# SẮP XẾP THEO THÀNH PHỐ
def sort_by_city():
    global df_raw, df_clean

    target_df = df_clean if df_clean is not None else df_raw

    if target_df is not None:
        target_df = target_df.sort_values(by="city").reset_index(drop=True)

        if df_clean is not None:
            df_clean = target_df

        else:
            df_raw = target_df

        show_table(target_df)
        status_label.config(text="Đã sắp xếp theo thành phố")

    else:
        messagebox.showwarning("Thông báo", "Vui lòng tải dữ liệu lên trước khi sắp xếp")



# TÌM KIẾM
def search():
    global df_raw, df_clean

    target_df = df_clean if df_clean is not None else df_raw

    if target_df is None:
        messagebox.showwarning("Chưa có dữ liệu", "Vui lòng chọn file trước")
        return
    
    search_window = tk.Toplevel(root)
    search_window.title("Tìm kiếm dữ liệu")
    search_window.geometry("400x200")
    search_window.resizable(False, False)

    # Tìm theo store_id
    frame_id = tk.LabelFrame(search_window, text="Tìm theo Store ID", padx=10, pady=10)
    frame_id.pack(fill="x", padx=10, pady=5)
    
    entry_id = tk.Entry(frame_id, width=30)
    entry_id.pack(side="left", padx=5)

    def search_by_storeid():
        store_id = entry_id.get().strip()
        if not store_id:
            messagebox.showwarning("Thiếu thông tin", "Vui lòng nhập Store ID")
            return
        
        result = target_df[target_df["store_id"].astype(str).str.upper() == store_id.upper()]
        
        if result.empty:
            messagebox.showinfo("Kết quả", f"Không tìm thấy store_id: {store_id}")
        else:
            show_table(result)
            status_label.config(text=f"Tìm thấy {len(result)} kết quả cho store_id: {store_id}")
            search_window.destroy()
    
    search_id_btn = tk.Button(frame_id, text="Tìm", command=search_by_storeid, width=10)
    search_id_btn.pack(side="left", padx=5)


    # Tìm theo city
    frame_city = tk.LabelFrame(search_window, text="Tìm theo City", padx=10, pady=10)
    frame_city.pack(fill="x", padx=10, pady=5)
    
    entry_city = tk.Entry(frame_city, width=30)
    entry_city.pack(side="left", padx=5)
    
    def search_by_city():
        city = entry_city.get().strip()
        if not city:
            messagebox.showwarning("Thiếu thông tin", "Vui lòng nhập City")
            return
        
        result = target_df[target_df["city"].str.upper() == city.upper()]
        
        if result.empty:
            messagebox.showinfo("Kết quả", f"Không tìm thấy city: {city}")
        else:
            show_table(result)
            status_label.config(text=f"Tìm thấy {len(result)} kết quả cho city: {city}")
            search_window.destroy()
    
    search_city_btn = tk.Button(frame_city, text="Tìm", command=search_by_city, width=10)
    search_city_btn.pack(side="left", padx=5)


    # Hiển thị lại dữ liệu ban đầu
    show_all_btn = tk.Button(
        search_window,
        text="Hiển thị tất cả",
        command=lambda: [show_table(target_df), status_label.config(text="Hiển thị toàn bộ dữ liệu"), search_window.destroy()],
        width=20
    )
    show_all_btn.pack(pady=10)



# XÓA DÒNG
def delete_row():
    global df_raw, df_clean
    
    target_df = df_clean if df_clean is not None else df_raw
    
    if target_df is None:
        messagebox.showwarning("Chưa có dữ liệu", "Vui lòng chọn file trước")
        return
    
    delete_window = tk.Toplevel(root)
    delete_window.title("Xóa dữ liệu")
    delete_window.geometry("400x170")
    delete_window.resizable(False, False)
    
    frame = tk.LabelFrame(delete_window, text="Xóa theo Store ID", padx=10, pady=10)
    frame.pack(fill="both", expand=True, padx=10, pady=10)
    
    tk.Label(frame, text="Nhập Store ID cần xóa:").pack(pady=5)
    
    entry_id = tk.Entry(frame, width=30)
    entry_id.pack(pady=5)

    def confirm_delete():
        store_id = entry_id.get().strip()
        if not store_id:
            messagebox.showwarning("Thiếu thông tin", "Vui lòng nhập Store ID")
            return

        matching_rows = target_df[target_df["store_id"].astype(str).str.upper() == store_id.upper()]
        
        if matching_rows.empty:
            messagebox.showinfo("Thông báo", f"Không tìm thấy store_id: {store_id}")
            return

        confirm = messagebox.askyesno("Xác nhận", f"Bạn có chắc muốn xóa {len(matching_rows)} dòng với store_id: {store_id}?")
        
        if confirm:
            if df_clean is not None:
                df_clean.drop(matching_rows.index, inplace=True)
                df_clean.reset_index(drop=True, inplace=True)
                show_table(df_clean)
            else:
                df_raw.drop(matching_rows.index, inplace=True)
                df_raw.reset_index(drop=True, inplace=True)
                show_table(df_raw)
            
            status_label.config(text=f"Đã xóa {len(matching_rows)} dòng với store_id: {store_id}")
            messagebox.showinfo("Hoàn tất", f"Đã xóa thành công!")
            delete_window.destroy()
    
    delete_btn = tk.Button(
        frame,
        text="Xóa",
        command=confirm_delete,
        width=15,
        cursor="hand2"
    )
    delete_btn.pack(pady=10)



# BIỂU ĐỒ
def show_charts():
    if df_clean is None:
        messagebox.showwarning("Chưa có dữ liệu", "Vui lòng làm sạch dữ liệu trước")
        return
    
    chart_window = tk.Toplevel(root)
    chart_window.title("Biểu đồ phân tích")
    chart_window.geometry("1350x700")

    fig, axes = plt.subplots(2, 3, figsize=(14, 8))
    fig.suptitle("Phân tích Highlands Coffee", fontsize=14, fontweight="bold")

    # 1. Biểu đồ cột: Doanh thu theo thành phố
    city_revenue = df_clean.groupby("city")["revenue"].sum().sort_values(ascending=False)
    axes[0, 0].bar(city_revenue.index, city_revenue.values, color=["#B42424", "#3AA81E", "#ECA438", "#31B1D8"])
    axes[0, 0].set_title("Doanh thu theo thành phố", fontsize=11)
    axes[0, 0].set_ylabel("Doanh thu")
    axes[0, 0].tick_params(axis="x", rotation=45, labelsize=9)
    
    # 2. Biểu đồ tròn: Tỉ lệ doanh thu
    axes[0, 1].pie(city_revenue.values, labels=city_revenue.index, autopct="%1.1f%%", startangle=90)
    axes[0, 1].set_title("Tỉ lệ doanh thu", fontsize=11)
    
    # 3. Biểu đồ đường: Lợi nhuận theo thành phố
    city_profit = df_clean.groupby("city")["profit"].sum().sort_values(ascending=False)
    axes[0, 2].plot(city_profit.index, city_profit.values, marker="o", color="green", linewidth=2)
    axes[0, 2].set_title("Lợi nhuận theo thành phố", fontsize=11)
    axes[0, 2].set_ylabel("Lợi nhuận")
    axes[0, 2].tick_params(axis="x", rotation=45, labelsize=9)
    axes[0, 2].grid(True, alpha=0.3)
    
    # 4. Biểu đồ cột ghép: Doanh thu và Lợi nhuận theo thành phố
    city_stats = df_clean.groupby("city")[["revenue", "profit"]].sum()
    x = np.arange(len(city_stats.index))
    width = 0.35
    axes[1, 0].bar(x - width / 2, city_stats["revenue"], width, label="Doanh thu", color="blue")
    axes[1, 0].bar(x + width / 2, city_stats["profit"], width, label="Lợi nhuận", color="orange")
    axes[1, 0].set_title("Doanh thu và Lợi nhuận theo thành phố")
    axes[1, 0].set_xlabel("Thành phố")
    axes[1, 0].set_xticks(x)
    axes[1, 0].set_xticklabels(city_stats.index, rotation=45)
    axes[1, 0].legend()
    
    # 5. Biểu đồ phân tán: Mối quan hệ giữa số khách hàng và doanh thu
    axes[1, 1].scatter(df_clean["customers"], df_clean["revenue"], alpha=0.6, color="purple")
    axes[1, 1].set_title("Mối quan hệ: Khách hàng với Doanh thu")
    axes[1, 1].set_xlabel("Số khách hàng")
    axes[1, 1].set_ylabel("Doanh thu (VND)")
    axes[1, 1].grid(True, alpha=0.3)
    
    # Thông tin tổng quan
    axes[1, 2].axis("off")

    total_revenue = df_clean["revenue"].sum()
    total_profit = df_clean["profit"].sum()
    total_customers = df_clean["customers"].sum()
    stats_text = "TỔNG QUAN\n\n"
    stats_text += f"Tổng doanh thu:\n{total_revenue:,.0f} VND\n\n"
    stats_text += f"Tổng lợi nhuận:\n{total_profit:,.0f} VND\n\n"
    stats_text += f"Tổng khách hàng:\n{total_customers:,}"
    axes[1, 2].text(
        0.5, 0.5,
        stats_text,
        ha="center",
        va="center", 
        fontsize=11,
        bbox=dict(boxstyle="round", facecolor="wheat", alpha=0.5)
    )
    
    plt.tight_layout()
    

    canvas = FigureCanvasTkAgg(fig, master=chart_window)
    canvas.draw()
    canvas.get_tk_widget().pack(fill=tk.BOTH, expand=True)


# PHÂN TÍCH & DỰ ĐOÁN
def analyze_and_predict():
    if df_clean is None:
        messagebox.showwarning("Chưa có dữ liệu", "Vui lòng làm sạch dữ liệu trước")
        return


    analysis_window = tk.Toplevel(root)
    analysis_window.title("Phân tích & Dự đoán")
    analysis_window.geometry("605x600")

    text_frame = tk.Frame(analysis_window)
    text_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)
    
    scrollbar = tk.Scrollbar(text_frame)
    scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
    
    text_widget = tk.Text(
        text_frame,
        wrap=tk.WORD,
        yscrollcommand=scrollbar.set, 
        font=("Courier New", 10)
    )
    text_widget.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
    scrollbar.config(command=text_widget.yview)
    

    analysis_text = "=" * 70 + "\n"
    analysis_text += "  PHÂN TÍCH KẾT QUẢ KINH DOANH\n"
    analysis_text += "=" * 70 + "\n\n"


    analysis_text += "TỔNG QUAN\n"
    analysis_text += "-" * 70 + "\n"
    total_revenue = df_clean["revenue"].sum()
    total_profit = df_clean["profit"].sum()
    total_customers = df_clean["customers"].sum()
    

    analysis_text += f"  Tổng doanh thu:        {total_revenue:>20,} VND\n"
    analysis_text += f"  Tổng lợi nhuận:        {total_profit:>20,} VND\n"
    analysis_text += f"  Tỉ lệ lợi nhuận:       {(total_profit / total_revenue * 100):>19.1f}%\n"
    analysis_text += f"  Tổng khách hàng:       {total_customers:>20,}\n"
    analysis_text += f"  Số cửa hàng:           {df_clean["store_name"].nunique():>20}\n\n"


    analysis_text += "  THEO THÀNH PHỐ\n"
    analysis_text += "-" * 70 + "\n"
    city_stats = df_clean.groupby("city").agg({
        "revenue": "sum",
        "profit": "sum",
        "customers": "sum"
    }).sort_values("revenue", ascending=False)
    
    for city, row in city_stats.iterrows():
        pct = row["revenue"] / total_revenue * 100
        analysis_text += f"\n  {city}:\n"
        analysis_text += f"    - Doanh thu:     {row["revenue"]:>15,} VND ({pct:.1f}%)\n"
        analysis_text += f"    - Lợi nhuận:     {row["profit"]:>15,} VND\n"
        analysis_text += f"    - Khách hàng:    {row["customers"]:>15,}\n"
    

    analysis_text += "\n3 CỬA HÀNG CÓ DOANH THU XUẤT SẮC NHẤT\n"
    analysis_text += "-" * 70 + "\n"
    top_3 = df_clean.nlargest(3, "revenue")[["store_name", "city", "revenue", "profit"]]
    for idx, (_, row) in enumerate(top_3.iterrows(), 1):
        analysis_text += f"\n  {idx}. {row["store_name"]} ({row["city"]})\n"
        analysis_text += f"     Doanh thu: {row["revenue"]:,} VND\n"
        analysis_text += f"     Lợi nhuận: {row["profit"]:,} VND\n"


    analysis_text += "\nDỰ ĐOÁN NĂM TỚI\n"
    analysis_text += "-" * 70 + "\n"

    df_temp = df_clean.copy()
    df_temp["month_dt"] = pd.to_datetime(df_temp["month"], format="%m-%Y", errors="coerce")
    monthly_data = df_temp.groupby("month_dt").agg({
        "revenue": "sum",
        "profit": "sum"
    }).sort_index()
    
    if len(monthly_data) >= 2:
        first_month_revenue = monthly_data["revenue"].iloc[0]
        last_month_revenue = monthly_data["revenue"].iloc[-1]
        
        if first_month_revenue > 0:
            growth_rate = ((last_month_revenue - first_month_revenue) / first_month_revenue) * 100

            predicted_revenue = total_revenue * (1 + growth_rate / 100)
            predicted_profit = total_profit * (1 + growth_rate / 100)
            
            analysis_text += f"  Tốc độ tăng trưởng:           {growth_rate:>10.1f}%\n\n"
            analysis_text += f"  Dự kiến doanh thu năm tới:    {predicted_revenue:>15,.0f} VND\n"
            analysis_text += f"  Dự kiến lợi nhuận năm tới:    {predicted_profit:>15,.0f} VND\n"
            analysis_text += f"\n  Tăng trưởng dự kiến:        {(predicted_revenue - total_revenue):>15,.0f} VND\n"
        else:
            analysis_text += "  Không đủ dữ liệu để tính tốc độ tăng trưởng\n"
    else:
        analysis_text += "  Cần ít nhất 2 tháng dữ liệu để dự đoán\n"


    analysis_text += "\nKHUYẾN NGHỊ\n"
    analysis_text += "-" * 70 + "\n"

    best_city = city_stats["revenue"].idxmax()
    analysis_text += f"  - Thành phố dẫn đầu: {best_city}\n"
    analysis_text += f"  - Nên đẩy mạnh hoạt động tại {best_city}\n\n"
    
    worst_3 = df_clean.nsmallest(3, "revenue")[["store_name", "city"]]
    analysis_text += "  - Cửa hàng cần cải thiện:\n"
    for _, row in worst_3.iterrows():
        analysis_text += f"    + {row["store_name"]} ({row["city"]})\n"
    
    analysis_text += "\n" + "=" * 70 + "\n"


    text_widget.insert("1.0", analysis_text)
    text_widget.config(state=tk.DISABLED)



# RESET
def reset():
    global df_raw, df_clean

    df_raw = None
    df_clean = None

    table.delete(*table.get_children())
    input_label.config(text="Chưa chọn file")
    status_label.config(text="")
    table_info.config(text="")

    clean_btn.config(state="disabled")
    save_btn.config(state="disabled")
    id_sort_btn.config(state="disabled")
    city_sort_btn.config(state="disabled")
    search_btn.config(state="disabled")
    delete_btn.config(state="disabled")
    analyze_btn.config(state="disabled")
    chart_btn.config(state="disabled")



# GUI
root = tk.Tk()
root.title("Trình xử lí dữ liệu Highlands Coffee")
root.geometry("1200x650")


    # Header
header_frame = tk.Frame(root, bg="#2F5496", height=60)
header_frame.pack(fill="x")
header_frame.pack_propagate(False)

title_label = tk.Label(
    header_frame,
    text="Trình xử lí dữ liệu Highlands Coffee",
    bg="#2F5496",
    fg="#FFFFFF",
    font=("Arial", 16, "bold")
)
title_label.pack(pady=15)


    # Chức năng - Control
control_frame = tk.Frame(root)
control_frame.pack(pady=10)


control_row1 = tk.Frame(control_frame)
control_row1.pack()

select_btn = tk.Button(
    control_row1,
    text="Chọn file CSV/Excel",
    command=select_file,
    width=20,
    font=("Arial", 10, "bold"),
    cursor="hand2"
)
select_btn.pack(side="left", padx=5)

clean_btn = tk.Button(
    control_row1,
    text="Làm sạch dữ liệu",
    command=clean_data,
    width=18,
    font=("Arial", 10, "bold"),
    cursor="hand2",
    state="disabled"
)
clean_btn.pack(side="left", padx=5)

id_sort_btn = tk.Button(
    control_row1,
    text="Sắp xếp theo ID cửa hàng",
    command=sort_by_storeid,
    width=25,
    font=("Arial", 10, "bold"),
    cursor="hand2",
    state="disabled"
)
id_sort_btn.pack(side="left", padx=5)

city_sort_btn = tk.Button(
    control_row1,
    text="Sắp xếp theo thành phố",
    command=sort_by_city,
    width=25,
    font=("Arial", 10, "bold"),
    cursor="hand2",
    state="disabled"
)
city_sort_btn.pack(side="left", padx=5)

search_btn = tk.Button(
    control_row1,
    text="Tìm kiếm",
    command=search,
    width=20,
    font=("Arial", 10, "bold"),
    cursor="hand2",
    state="disabled",
)
search_btn.pack(side="left", padx=5)

control_row2 = tk.Frame(control_frame)
control_row2.pack(pady=5)

save_btn = tk.Button(
    control_row2,
    text="Lưu file",
    command=save_file,
    width=18,
    font=("Arial", 10, "bold"),
    cursor="hand2",
    state="disabled"
)
save_btn.pack(side="left", padx=5)

reset_btn = tk.Button(
    control_row2,
    text="Reset",
    command=reset,
    width=12,
    font=("Arial", 10, "bold"),
    cursor="hand2"
)
reset_btn.pack(side="left", padx=5)

delete_btn = tk.Button(
    control_row2,
    text="Xóa dòng",
    command=delete_row,
    width=20,
    font=("Arial", 10, "bold"),
    cursor="hand2",
    state="disabled",
)
delete_btn.pack(side="left", padx=5)

chart_btn = tk.Button(
    control_row2,
    text="Biểu đồ",
    command=show_charts,
    width=20,
    font=("Arial", 10, "bold"),
    cursor="hand2",
    state="disabled"
)
chart_btn.pack(side="left", padx=5)

analyze_btn = tk.Button(
    control_row2,
    text="Phân tích & Dự đoán",
    command=analyze_and_predict,
    width=25,
    font=("Arial", 10, "bold"),
    cursor="hand2",
    state="disabled"
)
analyze_btn.pack(side="left", padx=5)


    # Thông tin - Info
info_frame = tk.Frame(root)
info_frame.pack(pady=5)

input_label = tk.Label(
    info_frame,
    text="Chưa chọn file",
    font=("Arial", 9)
)
input_label.pack()

status_label = tk.Label(
    info_frame,
    text="",
    font=("Arial", 9, "italic")
)
status_label.pack()


    # Bảng - Table
table_frame = tk.Frame(root)
table_frame.pack(fill="both", expand=True, padx=10, pady=5)

scroll_y = tk.Scrollbar(table_frame, orient="vertical")
scroll_x = tk.Scrollbar(table_frame, orient="horizontal")

table = ttk.Treeview(
    table_frame,
    yscrollcommand=scroll_y.set,
    xscrollcommand=scroll_x.set
)

scroll_y.config(command=table.yview)
scroll_x.config(command=table.xview)

scroll_y.pack(side="right", fill="y")
scroll_x.pack(side="bottom", fill="x")
table.pack(fill="both", expand=True)


    # Footer
footer_frame = tk.Frame(root)
footer_frame.pack(fill="x", pady=5)

table_info = tk.Label(footer_frame, text="", font=("Arial", 8))
table_info.pack()



# CHẠY GUI 
root.mainloop()